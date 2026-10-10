import csv
import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pdf2image import convert_from_bytes
from PIL import Image
from pypdf import PdfReader
from rapidfuzz.fuzz import ratio, token_set_ratio
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import (
    CanonicalProduct,
    Ingredient,
    InventoryBatch,
    PriceObservation,
    ProductCandidate,
    ProductIngredient,
    ProductSku,
    RawDocument,
    Warehouse,
    WarehouseBatchStock,
)
from app.models.enums import (
    ConfidenceLabel,
    ProcessingStatus,
    PublishStatus,
    RegulatoryStatus,
    RxOtcStatus,
)
from app.services.ai_provider import get_ai_provider
from app.services.fulfillment_service import FulfillmentRoutingService
from app.services.normalization import (
    extract_registration_number,
    normalize_for_match,
    normalize_registration_number,
    normalize_text,
)
from app.services.scoring import calculate_product_score


ALLOWED_SUFFIXES = {".csv", ".xlsx", ".json", ".pdf", ".png", ".jpg", ".jpeg"}

ALIASES: dict[str, list[str]] = {
    "name": [
        "name", "product_name", "ten_thuoc", "tên thuốc", "ten san pham", "tên sản phẩm", "ten", "tên"
    ],
    "registration_number": [
        "registration_number", "so_dang_ky", "số đăng ký", "sdk", "so_dk", "số đk", "ma_dang_ky"
    ],
    "manufacturer": [
        "manufacturer", "nha_san_xuat", "nhà sản xuất", "nsx", "cong_ty_san_xuat", "công ty sản xuất"
    ],
    "manufacturing_country": [
        "manufacturing_country", "nuoc_san_xuat", "nước sản xuất", "quoc_gia", "quốc gia", "country"
    ],
    "dosage_form": [
        "dosage_form", "dang_bao_che", "dạng bào chế", "bao_che", "bào chế"
    ],
    "package": [
        "package", "package_description", "quy_cach", "quy cách", "dong_goi", "đóng gói"
    ],
    "rx_otc": [
        "rx_otc", "rx_otc_status", "phan_loai", "phân loại", "loai_thuoc", "loại thuốc"
    ],
    "ingredients": [
        "ingredients", "active_ingredients", "hoat_chat", "hoạt chất", "thanh_phan", "thành phần"
    ],
    "strength": [
        "strength", "ham_luong", "hàm lượng", "nong_do", "nồng độ"
    ],
    "price": [
        "price", "gia_ban", "giá bán", "gia", "giá", "don_gia", "đơn giá", "selling_price"
    ],
    "image_url": [
        "image_url", "image", "hinh_anh", "hình ảnh", "anh", "ảnh", "anh_san_pham", "photo"
    ],
    "description": [
        "description", "mo_ta", "mô tả", "mo_ta_ngan", "mô tả ngắn", "chi_tiet", "chi tiết"
    ],
    "indications": [
        "indications", "cong_dung", "công dụng", "chi_dinh", "chỉ định", "tac_dung", "tác dụng"
    ],
    "usage_instructions": [
        "usage_instructions", "cach_dung", "cách dùng", "huong_dan_su_dung", "hướng dẫn sử dụng", "lieu_dung", "liều dùng"
    ],
    "contraindications": [
        "contraindications", "luu_y", "lưu ý", "chong_chi_dinh", "chống chỉ định", "canh_bao", "cảnh báo"
    ],
    "side_effects": [
        "side_effects", "tac_dung_phu", "tác dụng phụ"
    ],
    "storage_conditions": [
        "storage_conditions", "bao_quan", "bảo quản", "dieu_kien_bao_quan", "điều kiện bảo quản"
    ],
    "category": [
        "category", "danh_muc", "danh mục", "nhom_thuoc", "nhóm thuốc", "category_slug"
    ],
    "brand": [
        "brand", "thuong_hieu", "thương hiệu", "nhan_hang", "nhãn hàng"
    ],
    "batch_number": [
        "batch_number", "so_lo", "số lô", "ma_lo", "mã lô", "so_lo_san_xuat", "lot", "lot_number", "batch"
    ],
    "expiry_date": [
        "expiry_date", "han_dung", "hạn dùng", "han_su_dung", "hạn sử dụng", "hsd", "exp", "exp_date", "date_exp", "expire_date"
    ],
    "quantity": [
        "quantity", "so_luong", "số lượng", "sl", "sl_nhap", "so_luong_nhap", "stock", "ton_kho"
    ],
    "warehouse": [
        "warehouse", "kho", "kho_nhap", "chi_nhanh", "ma_kho", "mã kho", "warehouse_code"
    ],
}


def _parse_date(val: Any) -> date | None:
    if not val:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    try:
        if s.replace(".", "").isdigit():
            num = float(s)
            if 30000 < num < 60000:
                base = date(1899, 12, 30)
                return base + timedelta(days=int(num))
    except Exception:
        pass
    return None


def _parse_int(val: Any, default: int = 100) -> int:
    if val is None or val == "":
        return default
    try:
        s = str(val).strip().replace(",", "").replace(".", "")
        num = int(s)
        return num if num > 0 else default
    except Exception:
        return default


def generate_excel_template() -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách thuốc nhập kho"

    headers = [
        "Tên thuốc (*)",
        "Số đăng ký",
        "Số lô (*)",
        "Hạn sử dụng (*)",
        "Số lượng nhập (*)",
        "Giá bán (VNĐ)",
        "Quy cách đóng gói",
        "Dạng bào chế",
        "Hoạt chất chính",
        "Hàm lượng",
        "Nhà sản xuất",
        "Nước sản xuất",
        "Phân loại (OTC/RX)",
        "Công dụng / Chỉ định",
        "Cách dùng / Liều dùng",
    ]

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.row_dimensions[1].height = 28
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    sample_rows = [
        [
            "Panadol Extra Đỏ (Hộp 180 viên)",
            "VN-22013-19",
            "LOT-PND-202610",
            "2028-10-15",
            500,
            145000,
            "Hộp 15 vỉ x 12 viên",
            "Viên nén bao phim",
            "Paracetamol, Caffeine",
            "500mg, 65mg",
            "GlaxoSmithKline",
            "Việt Nam",
            "OTC",
            "Giảm đau nhanh: đau đầu, đau cơ, sốt",
            "Uống 1-2 viên mỗi 4-6 giờ khi cần",
        ],
        [
            "Augmentin 1g (Hộp 14 viên)",
            "VN-18234-14",
            "LOT-AUG-202609",
            "2028-09-30",
            200,
            265000,
            "Hộp 2 vỉ x 7 viên",
            "Viên nén bao phim",
            "Amoxicillin, Acid Clavulanic",
            "875mg, 125mg",
            "GlaxoSmithKline",
            "Pháp",
            "RX",
            "Kháng sinh điều trị nhiễm khuẩn đường hô hấp",
            "Dùng theo chỉ định của bác sĩ chuyên môn",
        ],
        [
            "Berberin 100mg Mộc Hoa Tràm (Lọ 100 viên)",
            "VD-24567-16",
            "LOT-BER-202611",
            "2029-05-20",
            1000,
            35000,
            "Lọ 100 viên",
            "Viên nén",
            "Berberin clorid",
            "100mg",
            "Dược phẩm OPC",
            "Việt Nam",
            "OTC",
            "Điều trị tiêu chảy, kiết lỵ, viêm đại tràng",
            "Uống 2-4 viên/lần, ngày 2 lần sau ăn",
        ],
        [
            "Efferalgan 500mg Sủi (Hộp 16 viên)",
            "VN-16543-13",
            "LOT-EFF-202610",
            "2028-11-01",
            800,
            68000,
            "Hộp 4 vỉ x 4 viên",
            "Viên sủi",
            "Paracetamol",
            "500mg",
            "UPSA SAS",
            "Pháp",
            "OTC",
            "Hạ sốt, giảm các cơn đau vừa và nhẹ",
            "Hòa tan 1 viên vào 200ml nước, uống khi đau",
        ],
    ]

    data_font = Font(name="Arial", size=10)
    for row_idx, row_data in enumerate(sample_rows, start=2):
        ws.row_dimensions[row_idx].height = 22
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in [3, 4, 5, 6, 13]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportValidationError(ValueError):
    pass


def safe_filename(filename: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]", "_", Path(filename).name)
    return stem[:180] or "upload.bin"


def validate_upload(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ImportValidationError(f"Định dạng {suffix or 'không xác định'} chưa được hỗ trợ")
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise ImportValidationError(f"Tệp vượt quá {settings.max_upload_mb} MB")
    if not content:
        raise ImportValidationError("Tệp rỗng")
    return suffix


def _parse_price(val: Any) -> Decimal | None:
    if val is None or val == "":
        return None
    try:
        if isinstance(val, (int, float)):
            return Decimal(str(val))
        s = str(val).strip()
        s = re.sub(r"[^\d.,]", "", s)
        if not s:
            return None
        if "." in s and "," in s:
            if s.rfind(",") > s.rfind("."):
                s = s.replace(".", "").replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "," in s:
            parts = s.split(",")
            if len(parts[-1]) == 3 and len(parts) > 1:
                s = s.replace(",", "")
            elif len(parts[-1]) <= 2:
                s = s.replace(",", ".")
            else:
                s = s.replace(",", "")
        elif "." in s:
            parts = s.split(".")
            if len(parts[-1]) == 3 and len(parts) > 1 and float(s) < 1000:
                s = s.replace(".", "")
        return Decimal(s)
    except Exception:
        return None


def _parse_ingredients(row: dict[str, Any]) -> list[dict[str, Any]]:
    ingredients_val = row.get("ingredients")
    strength_val = row.get("strength")
    if isinstance(ingredients_val, list):
        return ingredients_val
    if not ingredients_val:
        return []
    items: list[dict[str, Any]] = []
    raw_str = str(ingredients_val).strip()
    parts = [p.strip() for p in re.split(r"[,;+]", raw_str) if p.strip()]
    for p in parts:
        item: dict[str, Any] = {"name": p}
        if strength_val and len(parts) == 1:
            item["original_strength_text"] = str(strength_val).strip()
            match = re.search(r"(\d+(?:[.,]\d+)?)\s*([a-zA-Z%]+)", str(strength_val))
            if match:
                try:
                    item["strength_value"] = float(match.group(1).replace(",", "."))
                    item["strength_unit"] = match.group(2).lower()
                except Exception:
                    pass
        items.append(item)
    return items


def parse_rx_otc(val: Any) -> RxOtcStatus:
    if not val:
        return RxOtcStatus.OTC
    s = normalize_for_match(str(val))
    if any(k in s for k in ["khong ke don", "otc", "khong can don"]):
        return RxOtcStatus.OTC
    if any(k in s for k in ["ke don", "rx", "thuoc ke don", "can don"]):
        return RxOtcStatus.PRESCRIPTION
    return RxOtcStatus.OTC


def _canonical_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized_keys = {normalize_for_match(str(key)).replace(" ", "_"): value for key, value in row.items()}
    result: dict[str, Any] = {}
    for target, aliases in ALIASES.items():
        for alias in aliases:
            key = normalize_for_match(alias).replace(" ", "_")
            if key in normalized_keys and normalized_keys[key] not in (None, ""):
                result[target] = normalized_keys[key]
                break
    return result


def _rows_from_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [dict(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]


def _rows_from_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(item or "").strip() for item in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, values, strict=False)) for values in rows if any(value is not None for value in values)]


def _rows_from_json(content: bytes) -> list[dict[str, Any]]:
    data = json.loads(content.decode("utf-8"))
    if isinstance(data, dict):
        data = data.get("items", [data])
    if not isinstance(data, list) or not all(isinstance(item, dict) for item in data):
        raise ImportValidationError("JSON phải là object hoặc danh sách object")
    return data


def _extract_document_text(content: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            return text
        images = convert_from_bytes(content, dpi=200, first_page=1, last_page=min(5, len(reader.pages)))
        return "\n".join(_ocr_image(image) for image in images)
    return _ocr_image(Image.open(io.BytesIO(content)))


def _ocr_image(image: Image.Image) -> str:
    import pytesseract

    return pytesseract.image_to_string(image.convert("RGB"), lang="vie+eng")


def _row_from_document(text: str) -> dict[str, Any]:
    lines = [normalize_text(line) for line in text.splitlines() if normalize_text(line)]
    name = lines[0] if lines else "Tài liệu chưa xác định tên thuốc"
    manufacturer_match = re.search(r"(?:Nhà sản xuất|Sản xuất bởi)\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    ingredient_match = re.search(r"(?:Thành phần|Hoạt chất)\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    package_match = re.search(r"(?:Quy cách)\s*[:\-]\s*(.+)", text, re.IGNORECASE)
    return {
        "name": name,
        "registration_number": extract_registration_number(text),
        "manufacturer": manufacturer_match.group(1).strip() if manufacturer_match else None,
        "ingredients": ingredient_match.group(1).strip() if ingredient_match else None,
        "package": package_match.group(1).strip() if package_match else None,
    }


def parse_upload(filename: str, content: bytes) -> tuple[list[dict[str, Any]], str | None, dict[str, Any] | None]:
    suffix = validate_upload(filename, content)
    if suffix == ".csv":
        return _rows_from_csv(content), None, None
    if suffix == ".xlsx":
        return _rows_from_xlsx(content), None, None
    if suffix == ".json":
        return _rows_from_json(content), None, None
    text = _extract_document_text(content, suffix)
    result = get_ai_provider().extract_product(text)
    extracted = result.data.model_dump()
    extracted["name"] = extracted.pop("product_name", None)
    extracted["package"] = extracted.pop("package_description", None)
    extracted["rx_otc"] = extracted.pop("rx_otc_status", "UNKNOWN")
    return [extracted], text, {
        "provider": result.provider,
        "model": result.model,
        "prompt_hash": result.prompt_hash,
        "fallback_used": result.fallback_used,
        "warnings": result.warnings,
    }


def _find_matching_canonical(db: Session, row: dict[str, Any]) -> CanonicalProduct | None:
    raw_reg = str(row.get("registration_number") or "").strip()
    norm_reg = normalize_registration_number(raw_reg) if raw_reg else None
    if norm_reg:
        prod = db.scalar(
            select(CanonicalProduct).where(
                func.lower(CanonicalProduct.registration_number) == raw_reg.lower()
            )
        )
        if prod:
            return prod

    name = str(row.get("name") or "").strip()
    if not name:
        return None
    norm_name = normalize_for_match(name)

    # Exact name match (case-insensitive)
    prod = db.scalar(
        select(CanonicalProduct).where(
            func.lower(CanonicalProduct.canonical_name) == name.lower()
        )
    )
    if prod:
        return prod

    # Search candidates with first keyword
    words = [w for w in name.split() if len(w) > 2]
    first_word = words[0] if words else name[:5]
    potentials = db.scalars(
        select(CanonicalProduct).where(
            CanonicalProduct.canonical_name.ilike(f"%{first_word}%")
        ).limit(40)
    ).all()

    best_score = 0.0
    best_prod = None
    for item in potentials:
        item_norm = normalize_for_match(item.canonical_name)
        score = max(
            ratio(norm_name, item_norm),
            token_set_ratio(norm_name, item_norm),
        ) / 100.0
        if score > best_score:
            best_score = score
            best_prod = item

    if best_score >= 0.82 and best_prod is not None:
        return best_prod

    return None


def _allocate_batch_and_stock(
    db: Session,
    sku_id: int,
    product_id: int,
    wh: Warehouse | None,
    row: dict[str, Any],
    default_batch_number: str | None,
    default_expiry_date: str | None,
    default_quantity: int | None,
) -> tuple[str, str, int]:
    if not wh or not sku_id:
        return "", "", 0

    item_batch_no = (
        str(row.get("batch_number") or "").strip()
        or (default_batch_number or "").strip()
        or f"LOT-IMP-{product_id}-{date.today().strftime('%Y%m')}"
    )
    item_exp = (
        _parse_date(row.get("expiry_date"))
        or _parse_date(default_expiry_date)
        or (date.today() + timedelta(days=730))
    )
    item_qty = _parse_int(row.get("quantity"), default_quantity or 100)

    batch = db.scalar(
        select(InventoryBatch).where(
            InventoryBatch.sku_id == sku_id,
            InventoryBatch.batch_number == item_batch_no,
        )
    )
    if not batch:
        batch = InventoryBatch(
            sku_id=sku_id,
            batch_number=item_batch_no,
            expiry_date=item_exp,
            initial_quantity=item_qty,
            status="ACTIVE",
        )
        db.add(batch)
        db.flush()

    wh_stock = db.scalar(
        select(WarehouseBatchStock).where(
            WarehouseBatchStock.warehouse_id == wh.id,
            WarehouseBatchStock.batch_id == batch.id,
        )
    )
    if not wh_stock:
        wh_stock = WarehouseBatchStock(
            warehouse_id=wh.id,
            batch_id=batch.id,
            quantity_on_hand=item_qty,
            quantity_available=item_qty,
            quantity_reserved=0,
        )
        db.add(wh_stock)
    else:
        wh_stock.quantity_on_hand += item_qty
        wh_stock.quantity_available += item_qty

    return item_batch_no, item_exp.isoformat(), item_qty


def preview_import_rows(
    db: Session,
    filename: str,
    content: bytes,
    warehouse_id: int | None = None,
    default_batch_number: str | None = None,
    default_expiry_date: str | None = None,
    default_quantity: int | None = None,
) -> dict[str, Any]:
    rows, document_text, ai_metadata = parse_upload(filename, content)

    target_wh = None
    if warehouse_id:
        target_wh = db.get(Warehouse, warehouse_id)
    if not target_wh:
        target_wh = db.query(Warehouse).first()

    parsed_default_exp = _parse_date(default_expiry_date) or (date.today() + timedelta(days=730))
    def_batch = (default_batch_number or "").strip() or f"LOT-{date.today().strftime('%Y%m%d')}"
    def_qty = default_quantity if default_quantity and default_quantity > 0 else 100

    items_preview = []
    valid_count = 0
    duplicate_count = 0
    total_qty = 0

    for idx, raw_row in enumerate(rows, start=1):
        row = _canonical_row(raw_row)
        name = str(row.get("name") or "").strip()
        if not name:
            items_preview.append({
                "row": idx,
                "name": "(Thiếu tên thuốc trong tệp)",
                "registration_number": None,
                "batch_number": None,
                "expiry_date": None,
                "quantity": 0,
                "price": None,
                "manufacturer": None,
                "status": "INVALID",
                "status_label": "Lỗi: Thiếu tên",
                "is_near_expiry": False,
            })
            continue

        row_batch = str(row.get("batch_number") or "").strip() or def_batch
        row_exp = _parse_date(row.get("expiry_date")) or parsed_default_exp
        row_qty = _parse_int(row.get("quantity"), def_qty)
        price_val = _parse_price(row.get("price"))
        reg_num = str(row.get("registration_number") or "").strip() or None

        existing = _find_matching_canonical(db, row)
        is_dup = existing is not None

        if is_dup:
            duplicate_count += 1
            status_text = "EXISTING"
            status_label = "Thuốc đã có (Cập nhật lô & tồn kho)"
        else:
            status_text = "NEW"
            status_label = "Thuốc mới tạo vào kho"

        valid_count += 1
        total_qty += row_qty

        items_preview.append({
            "row": idx,
            "name": name,
            "registration_number": reg_num,
            "batch_number": row_batch,
            "expiry_date": row_exp.isoformat(),
            "quantity": row_qty,
            "price": float(price_val) if price_val is not None else None,
            "manufacturer": row.get("manufacturer"),
            "dosage_form": row.get("dosage_form"),
            "package": row.get("package"),
            "status": status_text,
            "status_label": status_label,
            "is_near_expiry": (row_exp - date.today()).days < 180,
        })

    return {
        "filename": filename,
        "warehouse_id": target_wh.id if target_wh else None,
        "warehouse_name": target_wh.name if target_wh else "Kho Mặc Định",
        "warehouse_code": target_wh.code if target_wh else "",
        "total_rows": len(rows),
        "valid_rows": valid_count,
        "duplicate_rows": duplicate_count,
        "invalid_rows": len(rows) - valid_count,
        "total_quantity": total_qty,
        "default_batch_applied": def_batch,
        "default_expiry_applied": parsed_default_exp.isoformat(),
        "items": items_preview,
    }


def import_rows(
    db: Session,
    source_id: int,
    filename: str,
    content: bytes,
    is_demo: bool = False,
    force_update: bool = False,
    auto_approve: bool = True,
    warehouse_id: int | None = None,
    default_batch_number: str | None = None,
    default_expiry_date: str | None = None,
    default_quantity: int | None = None,
) -> dict[str, Any]:
    rows, document_text, ai_metadata = parse_upload(filename, content)
    filename = safe_filename(filename)
    file_hash = hashlib.sha256(content).hexdigest()
    upload_path = settings.upload_storage / f"{file_hash[:16]}_{filename}"
    upload_path.parent.mkdir(parents=True, exist_ok=True)
    if not upload_path.exists():
        upload_path.write_bytes(content)

    created = 0
    updated = 0
    skipped = 0
    total_imported_quantity = 0
    errors: list[dict[str, Any]] = []
    processed_items: list[dict[str, Any]] = []

    target_wh = None
    if warehouse_id:
        target_wh = db.get(Warehouse, warehouse_id)
    if not target_wh:
        target_wh = db.query(Warehouse).first()

    for index, raw_row in enumerate(rows, start=1):
        row = _canonical_row(raw_row)
        if not row.get("name"):
            errors.append({"row": index, "error": "Thiếu tên thuốc"})
            continue

        row_json = json.dumps(raw_row, ensure_ascii=False, sort_keys=True, default=str)
        content_hash = hashlib.sha256(row_json.encode()).hexdigest()
        source_url = f"upload://{filename}#row={index}"

        existing_doc = db.scalar(
            select(RawDocument).where(
                RawDocument.source_id == source_id,
                RawDocument.source_url == source_url,
                RawDocument.content_hash == content_hash,
            )
        )

        existing_candidate = None
        if existing_doc:
            existing_candidate = db.scalar(
                select(ProductCandidate).where(
                    ProductCandidate.raw_document_id == existing_doc.id
                )
            )

        if existing_doc and not force_update and not auto_approve:
            skipped += 1
            processed_items.append({
                "row": index,
                "name": row["name"],
                "registration_number": row.get("registration_number"),
                "status": "SKIPPED",
                "reason": "Bản ghi trùng lặp",
            })
            continue

        if not existing_doc:
            document = RawDocument(
                source_id=source_id,
                source_url=source_url,
                content_type="application/json" if document_text is None else "text/plain",
                local_storage_path=str(upload_path),
                raw_text=document_text or row_json,
                content_hash=content_hash,
                http_status=None,
                file_size=len(content),
                metadata_json={"original_filename": filename, "row_number": index, "ai": ai_metadata},
                is_demo=is_demo,
            )
            db.add(document)
            db.flush()
        else:
            document = existing_doc

        ingredients = _parse_ingredients(row)
        price_val = _parse_price(row.get("price"))

        matched_prod = _find_matching_canonical(db, row)
        status_action = "EXTRACTED"
        prod: CanonicalProduct | None = None
        alloc_batch = ""
        alloc_exp = ""
        alloc_qty = 0

        if matched_prod:
            prod = matched_prod
            if force_update or auto_approve:
                if row.get("dosage_form"):
                    prod.dosage_form = str(row["dosage_form"])
                if row.get("package"):
                    prod.package_description = str(row["package"])
                if row.get("manufacturer"):
                    prod.manufacturer = str(row["manufacturer"])
                if row.get("manufacturing_country"):
                    prod.manufacturing_country = str(row["manufacturing_country"])
                if row.get("image_url"):
                    prod.image_url = str(row["image_url"])
                if row.get("description"):
                    prod.description = str(row["description"])
                if row.get("indications"):
                    prod.indications = str(row["indications"])
                if row.get("usage_instructions"):
                    prod.usage_instructions = str(row["usage_instructions"])
                if row.get("contraindications"):
                    prod.contraindications = str(row["contraindications"])
                if row.get("side_effects"):
                    prod.side_effects = str(row["side_effects"])
                if row.get("storage_conditions"):
                    prod.storage_conditions = str(row["storage_conditions"])
                if row.get("category"):
                    prod.category_slug = str(row["category"])
                if row.get("rx_otc"):
                    prod.rx_otc_status = parse_rx_otc(row["rx_otc"])
                if auto_approve:
                    prod.publish_status = PublishStatus.PUBLISHED

                if price_val is not None:
                    po = db.query(PriceObservation).filter(PriceObservation.product_id == prod.id).first()
                    if po:
                        po.observed_price = price_val
                        po.observed_at = datetime.now(timezone.utc)
                    else:
                        po = PriceObservation(
                            product_id=prod.id,
                            source_id=source_id,
                            observed_price=price_val,
                            currency="VND",
                            availability_text="Có sẵn",
                            source_url=source_url,
                        )
                        db.add(po)

                sku = FulfillmentRoutingService.get_or_create_default_sku(db, prod.id)
                if price_val is not None:
                    sku.base_price = price_val

                alloc_batch, alloc_exp, alloc_qty = _allocate_batch_and_stock(
                    db,
                    sku.id,
                    prod.id,
                    target_wh,
                    row,
                    default_batch_number,
                    default_expiry_date,
                    default_quantity,
                )
                total_imported_quantity += alloc_qty

                if auto_approve:
                    prod.publish_status = PublishStatus.PUBLISHED
                    prod.confidence_label = ConfidenceLabel.HIGH_OFFICIAL_MATCH
                    prod.overall_score = max(prod.overall_score or 0, 88)

                try:
                    from app.services.search_service import index_single_product
                    index_single_product(db, prod)
                except Exception:
                    pass

                updated += 1
                status_action = "UPDATED"
            else:
                updated += 1
                status_action = "MATCHED"

        elif auto_approve:
            rx_status = parse_rx_otc(row.get("rx_otc"))

            prod = CanonicalProduct(
                canonical_name=str(row["name"]),
                registration_number=str(row.get("registration_number") or "") or None,
                dosage_form=str(row.get("dosage_form") or "") or None,
                package_description=str(row.get("package") or "") or None,
                manufacturer=str(row.get("manufacturer") or "") or None,
                manufacturing_country=str(row.get("manufacturing_country") or "") or None,
                image_url=str(row.get("image_url") or "") or None,
                description=str(row.get("description") or "") or None,
                usage_instructions=str(row.get("usage_instructions") or "") or None,
                indications=str(row.get("indications") or "") or None,
                contraindications=str(row.get("contraindications") or "") or None,
                side_effects=str(row.get("side_effects") or "") or None,
                storage_conditions=str(row.get("storage_conditions") or "") or None,
                category_slug=str(row.get("category") or "") or None,
                regulatory_status=RegulatoryStatus.ACTIVE,
                rx_otc_status=rx_status,
                publish_status=PublishStatus.PUBLISHED,
                overall_score=85,
                confidence_label=ConfidenceLabel.HIGH_OFFICIAL_MATCH,
                is_demo=is_demo,
            )
            db.add(prod)
            db.flush()

            for ing in ingredients:
                ing_name = ing.get("name") if isinstance(ing, dict) else str(ing)
                if ing_name:
                    norm_ing = normalize_for_match(ing_name)
                    ing_entity = db.scalar(select(Ingredient).where(Ingredient.normalized_name == norm_ing))
                    if not ing_entity:
                        ing_entity = Ingredient(normalized_name=norm_ing, alternative_names=[ing_name])
                        db.add(ing_entity)
                        db.flush()
                    db.add(
                        ProductIngredient(
                            product_id=prod.id,
                            ingredient_id=ing_entity.id,
                            strength_value=ing.get("strength_value") if isinstance(ing, dict) else None,
                            strength_unit=ing.get("strength_unit") if isinstance(ing, dict) else None,
                            original_strength_text=ing.get("original_strength_text") if isinstance(ing, dict) else None,
                        )
                    )

            if price_val is not None:
                db.add(
                    PriceObservation(
                        product_id=prod.id,
                        source_id=source_id,
                        observed_price=price_val,
                        currency="VND",
                        availability_text="Có sẵn",
                        source_url=source_url,
                    )
                )

            sku = FulfillmentRoutingService.get_or_create_default_sku(db, prod.id)
            if price_val is not None:
                sku.base_price = price_val

            alloc_batch, alloc_exp, alloc_qty = _allocate_batch_and_stock(
                db,
                sku.id,
                prod.id,
                target_wh,
                row,
                default_batch_number,
                default_expiry_date,
                default_quantity,
            )
            total_imported_quantity += alloc_qty

            prod.publish_status = PublishStatus.PUBLISHED
            prod.confidence_label = ConfidenceLabel.HIGH_OFFICIAL_MATCH
            prod.overall_score = 90

            try:
                from app.services.search_service import index_single_product
                index_single_product(db, prod)
            except Exception:
                pass

            created += 1
            status_action = "CREATED"
        else:
            created += 1
            status_action = "EXTRACTED"

        present_fields = sum(bool(row.get(key)) for key in ("name", "registration_number", "manufacturer", "dosage_form", "package", "ingredients"))
        confidence = round(present_fields / 6, 2)
        target_status = ProcessingStatus.MATCHED if prod else ProcessingStatus.EXTRACTED

        if existing_candidate:
            existing_candidate.observed_name = str(row["name"])
            existing_candidate.normalized_name = normalize_for_match(str(row["name"]))
            existing_candidate.registration_number_text = str(row.get("registration_number") or "") or None
            existing_candidate.manufacturer_text = str(row.get("manufacturer") or "") or None
            existing_candidate.dosage_form_text = str(row.get("dosage_form") or "") or None
            existing_candidate.package_text = str(row.get("package") or "") or None
            existing_candidate.rx_otc_text = str(row.get("rx_otc") or "") or None
            existing_candidate.ingredients_json = ingredients
            existing_candidate.image_url = str(row.get("image_url") or "") or None
            existing_candidate.description = str(row.get("description") or "") or None
            existing_candidate.usage_instructions = str(row.get("usage_instructions") or "") or None
            existing_candidate.indications = str(row.get("indications") or "") or None
            existing_candidate.contraindications = str(row.get("contraindications") or "") or None
            existing_candidate.side_effects = str(row.get("side_effects") or "") or None
            existing_candidate.storage_conditions = str(row.get("storage_conditions") or "") or None
            existing_candidate.extraction_confidence = confidence
            if prod:
                existing_candidate.canonical_product_id = prod.id
                existing_candidate.processing_status = target_status
        else:
            candidate = ProductCandidate(
                raw_document_id=document.id,
                canonical_product_id=prod.id if prod else None,
                observed_name=str(row["name"]),
                normalized_name=normalize_for_match(str(row["name"])),
                registration_number_text=str(row.get("registration_number") or "") or None,
                manufacturer_text=str(row.get("manufacturer") or "") or None,
                dosage_form_text=str(row.get("dosage_form") or "") or None,
                package_text=str(row.get("package") or "") or None,
                rx_otc_text=str(row.get("rx_otc") or "") or None,
                ingredients_json=ingredients,
                image_url=str(row.get("image_url") or "") or None,
                description=str(row.get("description") or "") or None,
                usage_instructions=str(row.get("usage_instructions") or "") or None,
                indications=str(row.get("indications") or "") or None,
                contraindications=str(row.get("contraindications") or "") or None,
                side_effects=str(row.get("side_effects") or "") or None,
                storage_conditions=str(row.get("storage_conditions") or "") or None,
                extraction_confidence=confidence,
                extraction_method=(ai_metadata or {}).get("provider", "PARSER").upper() if document_text is not None else "PARSER",
                extraction_model=(ai_metadata or {}).get("model"),
                processing_status=target_status,
            )
            db.add(candidate)

        processed_items.append({
            "row": index,
            "name": row["name"],
            "registration_number": row.get("registration_number"),
            "price": float(price_val) if price_val else None,
            "manufacturer": row.get("manufacturer"),
            "batch_number": alloc_batch,
            "expiry_date": alloc_exp,
            "quantity": alloc_qty,
            "status": status_action,
            "canonical_product_id": prod.id if prod else None,
        })

    db.flush()
    return {
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "total_rows": len(rows),
        "total_quantity": total_imported_quantity,
        "warehouse_id": target_wh.id if target_wh else None,
        "warehouse_name": target_wh.name if target_wh else "",
        "items": processed_items,
    }
